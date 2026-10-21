"""Merge returned gold-set forms (data/gold_forms/*.docx, filled in by volunteers)
into data/gold_normalisation.tsv, ready for `python experiments/b0_gold.py`.

For every word row in a returned form:
  historical  the 1880s spelling printed in the form
  modern      what the volunteer wrote (rows left blank are skipped)
  checked_by  the name on the form (falls back to the file name)
  note        the volunteer's comment (name / English / OCR error / not sure / ...)
Rows already in the TSV for the same word are replaced, so re-running is safe.
Words the volunteer marked as an OCR error with no modern spelling are kept out.

Run:  python scripts/collect_gold_forms.py path/to/returned_forms/ [--dry-run]
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src import config  # noqa: E402

GOLD = config.ROOT / "data" / "gold_normalisation.tsv"
CANDIDATES = config.ROOT / "data" / "gold_candidates.tsv"
SAME_WORDS = {"same", "unchanged", "-", "\u2013", "\u2014", "ditto", "nochange"}


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace(" ", " ")).strip()


def read_form(path: Path) -> tuple[str, list[dict]]:
    import docx

    d = docx.Document(str(path))
    name = ""
    for para in d.paragraphs:
        m = re.match(r"\s*Your name:\s*(.*?)(Is isiXhosa your home language\?.*)?$", para.text)
        if m:
            name = clean(m.group(1).replace("_", ""))
            break
    rows = []
    for table in d.tables:
        for r in table.rows[1:]:  # skip header row
            cells = [clean(c.text) for c in r.cells]
            if len(cells) < 4 or not cells[0].isdigit():
                continue  # examples table ("e.g.") or anything unexpected
            num, hist, modern, comment = cells[:4]
            rows.append({"n": int(num), "historical": hist, "modern": modern, "note": comment})
    return name, rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", type=Path, help="folder holding the returned .docx forms")
    ap.add_argument("--dry-run", action="store_true", help="print what would be merged, write nothing")
    args = ap.parse_args()

    cand = set(pd.read_csv(CANDIDATES, sep="\t", dtype=str).historical.str.lower())
    new = []
    for f in sorted(args.folder.glob("*.docx")):
        if f.name.startswith("~$"):  # Word lock files
            continue
        name, rows = read_form(f)
        checker = name or f.stem
        filled = [r for r in rows if r["modern"]]
        ocr = [r for r in rows if not r["modern"] and "ocr" in r["note"].lower()]
        print(f"{f.name}: checker={checker!r}, {len(filled)}/{len(rows)} words filled, {len(ocr)} marked OCR error")
        for r in filled:
            if r["historical"].lower() not in cand:
                print(f"  [warn] {r['historical']!r} is not in gold_candidates.tsv; kept anyway")
            modern = r["modern"].split()[0].lower().strip(".,;")
            if modern in SAME_WORDS:  # "same", "unchanged", "-": the spelling has not changed
                modern = r["historical"].lower()
            new.append({"historical": r["historical"].lower(), "modern": modern,
                        "checked_by": checker, "note": r["note"]})
    if not new:
        raise SystemExit("no filled-in rows found")
    new = pd.DataFrame(new).drop_duplicates("historical", keep="last")

    old = pd.read_csv(GOLD, sep="\t", dtype=str).fillna("")
    merged = pd.concat([old[~old.historical.str.lower().isin(new.historical)], new], ignore_index=True)
    print(f"{len(new)} rows from forms; gold file goes from {len(old)} to {len(merged)} rows")
    if args.dry_run:
        print(new.head(10).to_string(index=False))
        return
    merged.to_csv(GOLD, sep="\t", index=False)
    print("wrote", GOLD, "- now run: python experiments/b0_gold.py")


if __name__ == "__main__":
    main()
